import os
import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
pptx_path = r"C:\Users\ASUS\.gemini\antigravity\scratch\mppt_simulator\MPPT_Teqdimat_20_Slayd_Yenilenmis.pptx"
prs = Presentation(pptx_path)

print(f"Total slides in presentation: {len(prs.slides)}")
assert len(prs.slides) == 20, f"Expected 20 slides, got {len(prs.slides)}"

for i, slide in enumerate(prs.slides):
    texts = []
    pics = 0
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text.strip():
                    texts.append(p.text.strip())
        if shape.shape_type == 13: # Picture
            pics += 1
    
    first_title = texts[1] if len(texts) > 1 else (texts[0] if texts else "NO TEXT")
    print(f"Slide {i+1:02d}: Pics={pics}, TextBlocks={len(texts)} | Title: {first_title[:45]}")

print("\n--- Verifying Slide 1 Details ---")
s1_texts = [p.text for s in prs.slides[0].shapes if s.has_text_frame for p in s.text_frame.paragraphs if p.text.strip()]
print("Slide 1 content:")
for t in s1_texts:
    print(f"  -> {t}")

assert any("Azərbaycan Texniki Universiteti" == t.strip() for t in s1_texts), "Azərbaycan Texniki Universiteti not found in slide 1!"
assert any("Həmidov Xalid" in t for t in s1_texts), "Həmidov Xalid not found in slide 1!"
print("\n[ALL CHECKS PASSED] Presentation structure is verified and robust.")
